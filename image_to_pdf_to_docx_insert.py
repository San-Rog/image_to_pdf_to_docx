import streamlit as st
import time
import re
import os
import io
import math
import papersize
from papersize import parse_papersize
from decimal import Decimal
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageOps, ImageDraw, ImageFont
from streamlit_extras.scroll_to_element import *

@st.cache_data     
def convertSize(tam):
    var = ['KB', 'MB', 'GB', 'TB', 'PB']
    size = 1024
    rest = tam
    for v, vr in enumerate(var):
        x = divmod(tam, size)
        tam = x[0]
        rest = x[1]
        if x[0] < 1000:
            break
    valFloat = f'{float(tam + rest/size):.2f}'
    valStr = f'{valFloat.replace('.', ',')}{vr}'.strip()
    return valStr.rjust(8)

@st.dialog(title=":red[Informações sobre o app] :material/deployed_code:", width="medium", 
           icon=":material/info:", on_dismiss="ignore") 
def mensInfoAll():
    dictButtons = {"botão": optAll, "funcionalidade": optText, "abrangência": scopeText}
    dfButtons = pd.DataFrame(dictButtons)
    st.table(dfButtons, border=True, width="stretch", height="stretch")
        
def changeCheckbox(num):
    if num == 0:
        if st.session_state['checkYes']:
            st.session_state['checkNo'] = False        
    else:
        if st.session_state['checkNo']:
            st.session_state['checkYes'] = False

@st.dialog(title=":red[Configuração de imagens] :material/create_new_folder:", 
           width="medium", icon=":material/settings:", on_dismiss="ignore") 
def mensCreateAll(fileUps):
    roleLabel = ["imagens com título sobreposto", "imagens sem título sobreposto"]
    intervUps = range(len(fileUps))
    itens = [w+1 for w in intervUps]
    itensBruts = [w for w in intervUps]
    names, types, sizes = ([file.name for file in fileUps], [file.type for file in fileUps], 
                           [convertSize(file.size) for file in fileUps])
    dictVals = {"item": itens, "nome": names, "tipo": types, "tamanho": sizes}
    keyVals = list(dictVals.keys())
    df = pd.DataFrame(dictVals)
    st.dataframe(
            df,
            column_config={
            keyVals[0]: st.column_config.TextColumn(width="content", alignment="left"), 
            keyVals[1]: st.column_config.TextColumn(width="content", alignment="left"), 
            keyVals[2]: st.column_config.TextColumn(width="content", alignment="left"), 
            keyVals[3]: st.column_config.TextColumn(width="content")},
            hide_index=True,
            width="stretch", 
            on_select="ignore"
    )
    colAngleTime, colOthers = st.columns(spec=2, vertical_alignment="center", width="stretch", 
                                         gap="medium")
    with colAngleTime:
        colBadAngle, colAngle = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                           gap="small")
        colBadAngle.badge(":material/screen_rotation_up:", width="stretch", help="Define o ângulo da imagem.", 
                          color="green") 
        colAngle.select_slider(label="Selecione o ângulo de rotação", options=optAllAngles, 
                               key="angleRotatedAll", value=optAllAngles[4], label_visibility="collapsed")
        colBadTime, colTime = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                         gap="small")
        colBadTime.badge(":material/timer:", width="stretch", help="Define o tempo entre os slides.", 
                         color="blue")
        colTime.slider(label="Selecione o tempo enre os slides", min_value=0.0, max_value=10.0, 
                       key="numSlidesAll", value=0.0, label_visibility="collapsed", step=0.01)
        colBadResol, colResol = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                           gap="small")
        colBadResol.badge(":material/badge:", width="stretch", help="Define a resolução da imagem.", 
                          color="yellow")
        colResol.slider(label="Selecione a resolução da imagem", min_value=200, max_value=1600, 
                       key="numResolAll", value=200, label_visibility="collapsed", step=1)
        colBadYes, colYes = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                       gap="small")
        colBadYes.badge(":material/position_top_right:", width="stretch", help="Se marcado, o título aparecerá sobreposto à imagem.", 
                        color="yellow")
        colYes.checkbox(label=roleLabel[0], key="checkYes", width="stretch", on_change=changeCheckbox, 
                        args=(0, )) 
    with colOthers:
        papers = list(optAllPapers.keys())
        margins = list(optAllMargins.keys())        
        colBadPapers, colPapers = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                             gap="small")
        colBadPapers.badge(":material/assignment:", width="stretch", help="Define o formato da página.", 
                           color="green")
        colPapers.selectbox(label="Selecione o formato do papel", options=papers, label_visibility="collapsed", 
                            key="papelSelAll", width="stretch", index=6)
        colBadOrients, colOrients = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                             gap="small")
        colBadOrients.badge(":material/landscape:", width="stretch", help="Define a orientação da página.", 
                            color="blue")
        colOrients.selectbox(label="Selecione a orientação do papel", options=optAllOrients, label_visibility="collapsed", 
                             key="orientSelAll", width="stretch", index=None)
        
        colBadMargins, colMargins = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                               gap="small")
        colBadMargins.badge(":material/margin:", width="stretch", help="Define a margem da página.", 
                            color="yellow")
        colMargins.selectbox(label="Selecione o formato do papel", options=margins, label_visibility="collapsed", 
                             key="marginSelAll", width="stretch", index=4)
        colBadNo, colNo = st.columns([0.8, 10], vertical_alignment="center", width="stretch", 
                                           gap="small")
        colBadNo.badge(":material/position_bottom_right:", width="stretch", help="Se marcado, o título não aparecerá. Também não aparecerá se esta e a caixa à esquerda ficarem vazias.", 
                        color="yellow")
        colNo.checkbox(label=roleLabel[1], key="checkNo", width="stretch", on_change=changeCheckbox, 
                       args=(1, ))     
    
@st.dialog(title=":blue[**Rotação e exibição de imagem**]", width='large', icon=':material/360:', 
           on_dismiss='ignore')
def mensSliderUnique(*args):
    imgRotated = args[0]
    nameAll = args[1]
    angleOld = st.session_state['bytesAll'][nameAll][0][3]
    nameSplit = nameAll.split('_')
    name, ext = os.path.splitext(nameSplit[0])
    if ext.strip() != '':
        nameFile = nameSplit[0]
    else:
        nameFile = '_'.join(nameSplit[:2])    
    pos = math.floor(len(optAllAngles)/2)
    angleSel = st.select_slider(label="Selecione o ângulo de rotação", options=optAllAngles, 
                                key="angleRotated", label_visibility="collapsed")
    if angleSel != 0: 
        rotatedImg = imgRotated.rotate(angleSel, expand=True)
        name, ext = os.path.splitext(nameFile)
        nameRotated = f'{name}_{angleSel}_graus{ext}'
    else:
        rotatedImg = imgRotated
        nameRotated = nameFile
    if angleSel != angleOld:
        st.session_state['bytesAll'][nameAll][0][3] = angleSel
    with st.container(border=False, vertical_alignment="top", horizontal_alignment="center"):
        st.image(rotatedImg, caption=nameRotated)

@st.dialog(title=":green[**Exibição de imagens como slides**]", width='medium', icon=":material/wallpaper_slideshow:")
def mensSliderAll():
    keysData = list(st.session_state['bytesAll'].keys())
    timeSleep = st.session_state['numSlidesAll']
    placeImg = st.empty()
    nBts = len(keysData)
    msg = st.toast(f"Exibindo imagens como slides em {timeSleep}s...", icon=":material/wallpaper_slideshow:", duration="short")
    for k, key in enumerate(keysData):
        bytesData = st.session_state['bytesAll'][key]
        for bt, btData in enumerate(bytesData):
            nameFile = btData[0]
            bytesDataRed = btData[1]
            angleFile = btData[3]
            textSld = f"● {nameFile} ● {timeSleep}s ● {angleFile}°"
            msg.toast(textSld, icon=":material/wallpaper_slideshow:", duration="short")
            placeImg.image(bytesDataRed, caption=nameFile)
            time.sleep(timeSleep)

@st.dialog(title=":red[**Janela de download - arquivo Docx**]", width="small", icon=":material/docs:")
def mensDownDocx(*args):
    sizeCols = args[-1]
    colSucc, colDown = st.columns(sizeCols, vertical_alignment="center", width="stretch")
    colSucc.success("Operação finalizada com sucesso!", icon=":material/done_all:", width="stretch")
    if colDown.download_button(
                key="downDocx", 
                label=args[0],
                data=args[1],
                file_name=args[2],
                mime=args[3], 
                icon=":material/download_2:"
    ):
        st.rerun()
        
@st.dialog(title=":blue[**Janela de download - arquivo Pdf**]", width="small", icon=":material/picture_as_pdf:")
def mensDownPdf(*args):
    imagens = args[0]
    resol = args[1]
    pdfMerge = args[2]
    sizeCols = args[3]
    imagens[0].save(
            pdfMerge, 
            "PDF", 
            resolution=float(resol), 
            save_all=True, 
            append_images=imagens[1:]
    )
    colSucc, colDown = st.columns(sizeCols, vertical_alignment="center", width="stretch")
    colSucc.success("Operação finalizada com sucesso!", icon=":material/done_all:", width="stretch")
    with open(pdfMerge, "rb") as f:
        down = colDown.download_button(
            key="downPdf",
            label="",
            data=f,
            file_name="imagens_reunidas.pdf",
            mime="application/pdf", 
            icon=":material/download_2:"            
        )   
    if down: 
        st.rerun()
        
@st.dialog(title=":red[**Alerta sobre a funcionalidade**]", width="medium", icon=":material/warning:")
def mensAlert(mensText):
    st.warning(mensText, width="stretch")
    
@st.dialog(title=":red[**Erro no funcionamento do aplicativo**]", width="medium", icon=":material/warning:")
def mensError(mensError):
    st.error(mensError, width="stretch")

@st.cache_data   
def fullFiles(uploadedFiles, symbol):
    angleSel = st.session_state['angleRotatedAll']
    bytesAllDatas = {}
    keys = list(buttSymbs.keys())
    key = buttSymbs[keys[0]][-1]
    for up, uploadedFile in enumerate(uploadedFiles):
        nameFile = uploadedFile.name
        img = Image.open(uploadedFile)
        imgRotated = img.rotate(angleSel, expand=True)
        bytesData = imgRotated    
        upStr = str(up+1).zfill(6) 
        keyFile = f'{nameFile}_{upStr}_{key}'
        bytesAllDatas.setdefault(keyFile, [])
        bytesAllDatas[keyFile].append([nameFile, bytesData, keyFile, angleSel, symbol, uploadedFile])
    return bytesAllDatas

@st.cache_data   
def defineVector(categ, adic=None):
    paperFiles = st.session_state['papelSelAll']
    orientFiles = st.session_state['orientSelAll']
    marginFiles = optAllMargins[st.session_state['marginSelAll']]
    resolFiles = st.session_state['numResolAll']
    if categ == 0:
        uploades = st.session_state['allUpLoads']
        angleFiles = st.session_state['angleRotatedAll']
        keysData = list(st.session_state['bytesAll'].keys())
        return(uploades, angleFiles, resolFiles, paperFiles, orientFiles, marginFiles, keysData) 
    elif categ == 1:
        sizes = parse_papersize(paperFiles, "in")
        if orientFiles == optAllOrients[1]:
            sizeOne = float(Decimal(sizes[0]))
            sizeTwo = float(Decimal(sizes[1]))
        else:
            sizeOne = float(Decimal(sizes[1]))
            sizeTwo = float(Decimal(sizes[0]))
        widthPaper = int(sizeOne*resolFiles)
        heightPaper = int(sizeTwo*resolFiles)
        sizePaper = (widthPaper, heightPaper)
        newWidth = int(widthPaper*marginFiles)
        try:
            font = ImageFont.truetype("arial.ttf", size=15)
        except IOError:
            font = ImageFont.load_default()
        return(sizePaper, newWidth, font, widthPaper, heightPaper)
    
@st.cache_data        
def modifyImage(upload, newWidth, angleFiles):
    img = Image.open(upload)
    img = img.rotate(angleFiles, expand=True)
    width, height = img.size
    prop = newWidth/float(width)
    newHeight = int(float(height)*prop)
    img = img.resize((newWidth, newHeight), Image.Resampling.LANCZOS)
    return img
    
def saveMultImgPdf():
    uploades, angleFiles, resolFiles, paperFiles, orientFiles, marginFiles, keysData = defineVector(0)
    sizePaper, newWidth, font, widthPaper, heightPaper = defineVector(1)
    imagens = []
    pdfMerge = "documento_final.pdf"
    nBts = len(uploades)
    for u, upload in enumerate(uploades):
        bytesData = st.session_state['bytesAll'][keysData[u]][0]
        nameFile = bytesData[0]
        angleStr = bytesData[3]
        funcFile = bytesData[4] 
        newName = designateImgs(nameFile, u, nBts, angleStr, funcFile, 1)
        img = modifyImage(upload, newWidth, angleStr)
        canvas = Image.new("RGB", sizePaper, (255, 255, 255))
        drawImage = ImageDraw.Draw(canvas)
        posX = (widthPaper - img.width)//2
        posY = (heightPaper - img.height)//2
        posXtext = widthPaper//2
        if st.session_state['checkYes']:
            drawImage.text((posXtext, posY-50), newName, fill="blue", font=font)
        canvas.paste(img, (posX, posY))
        imagens.append(canvas)
    mensDownPdf(imagens, resolFiles, pdfMerge, [12, 1.5])
    
def saveMultImgDocx():
    uploades, angleFiles, resolFiles, paperFiles, orientFiles, marginFiles, keysData = defineVector(0)
    sizePaper, newWidth, font, widthPaper, heightPaper = defineVector(1)
    nBts = len(uploades)
    nLim = nBts - 1
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(sizePaper[0])
    section.page_height = Inches(sizePaper[1])
    nBts = len(uploades)
    for u, upload in enumerate(uploades):
        bytesData = st.session_state['bytesAll'][keysData[u]][0]
        nameFile = bytesData[0]
        angleStr = bytesData[3]
        funcFile = bytesData[4] 
        imgLoad = bytesData[-1]
        img = Image.open(imgLoad)
        imRotated = modifyImage(upload, newWidth, angleStr)
        imgBytes = io.BytesIO()
        imRotated.save(imgBytes, format='PNG')
        imgBytes.seek(0)
        newSize = int(widthPaper/resolFiles)
        newMargin = newSize*marginFiles
        newCaption = designateImgs(nameFile, u, nBts, angleStr, funcFile, 1)
        if st.session_state['checkYes']:
            titulo = doc.add_paragraph(newCaption)
            titulo.style = 'Caption'  
            formato_fonte = titulo.runs[0].font
            formato_fonte.size = Pt(8)  
        doc.add_picture(imgBytes, width=Inches(newMargin))
        try:
            lastAnte = doc.paragraphs[-2]
            lastAnte.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except:
            pass
        lastPara = doc.paragraphs[-1]
        lastPara.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if u < nLim:
            doc.add_page_break()
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    mensDownDocx("", buffer.getvalue(), "imagens_reunidas.docx", 
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 
                 [12, 1.5])
    
def operationFiles(*args):
    mode = args[0]
    if mode in [0, 4]:
        job = args[1]
        keys = list(buttSymbs.keys())
        key = buttSymbs[keys[0]][-1]
        keyFile = job.split(key)[0] + key
        item = list(st.session_state['bytesAll'].keys()).index(keyFile)
        upLoad = st.session_state['allUpLoads'][item]
        if mode != 0:
            image = Image.open(upLoad)
            st.session_state['angleRotated'] = 0
            mensSliderUnique(image, keyFile)
            st.session_state['clickAngle'] = True

def zeraVal():
   st.session_state['numSldImg'] = 0
   st.session_state['containers'] = False
   st.session_state['disabPillTwo'] = True

def fromToImg():
    value = st.session_state['numSldImg']
    if value > 0: 
        value -= 1
        valMaius = value + 1
        valMinus = value - 1
        st.session_state['disabPillTwo'] = False
    else:
        st.session_state['disabPillTwo'] = True
    return 

@st.cache_data   
def designateImgs(nameFile, bt, nBts, angle, funcFile, categ):
    if categ == 0:
        textImg = f'**{bt+1}/ {nBts}** ●  :blue[**{nameFile}**]  ●  :blue[**{angle}º**]'
        textImg += f'● :blue[**{funcFile}**]'
    else:
        textImg = f'# {bt+1}/{nBts} - {nameFile} - {angle}º'
    return textImg

def backForward(*args):
    num = args[0]
    jobSplit = args[1]
    elem = args[2]
    fileBytes = st.session_state['bytesAll']
    nUps = len(fileBytes)
    listFiles = list(fileBytes.keys())
    buttClick = '_'.join(jobSplit[:-3])
    ind = listFiles.index(buttClick)
    if elem == 0:
        numPos = num - 1
        if numPos <= 0:
            numPos = 0
    else: 
        numPos = num + 1
        if numPos >= nUps:
            numPos = nUps - 1
    scroll_to_element(listFiles[numPos])
    
def commandButt(job, sufix, bytesData=None, num=None):
    sep = '_'
    jobSplit = job.split(sep)
    suf = f'{jobSplit[-3]}{sep}{jobSplit[-2]}'
    elem = sufix.index(suf)
    match elem:
        case 0 | 3: 
            backForward(num, jobSplit, elem) 
        case 1: 
            scroll_to_element("keyPillFive")
            scroll_to_element(st.session_state['keyFirst'])
        case 2: 
            scroll_to_element(st.session_state['keyLast']) 
        case 4: 
            operationFiles(elem, job)
    if elem != 0:
        st.session_state['numSldImg'] = 0
    return

@st.cache_data   
def orderFile(upLoads, categ):
    if categ in [1, 2]:
        item = lambda up: up.name
    else:
        item = lambda up: f'{up.name}_{up.size}'
    elemsLoad = sorted([(item(upLoad)) for upLoad in upLoads])
    newUpLoad = []
    for elem in elemsLoad:
        for upLoad in upLoads:
            elemUp = item(upLoad)
            if elem == elemUp:
                if upLoad not in newUpLoad:
                    newUpLoad.append(upLoad)
    if categ in [1, 3]:
        newUpLoad = newUpLoad
    else:
        newUpLoad = list(reversed(newUpLoad))
    return newUpLoad
    
def changePill(uploadedFiles, mode):
    if mode == 0:
        valKeyPill = st.session_state['keyPill']
        optSel = optFiles.index(valKeyPill)
        match optSel:
            case 0:
                st.session_state['containers'] = True
                st.session_state['allUpLoads'] = uploadedFiles
                st.session_state['bytesAll'] = fullFiles(uploadedFiles, valKeyPill)
            case 1 | 2 | 3 | 4:
                st.session_state['containers'] = True
                uploadedFiles = orderFile(uploadedFiles, optSel)
                st.session_state['allUpLoads'] = uploadedFiles
                st.session_state['bytesAll'] = fullFiles(uploadedFiles, valKeyPill)
            case 5:
                timeSleep = st.session_state['numSlidesAll']
                if timeSleep == 0:
                    mensText = f":blue[**:material/timer_play:**] Tempo de exibição dos slides igual a :blue[**{int(timeSleep)}s**]. " \
                               "Altere esse parâmetro usando o botão :blue[**:material/factory:**]."
                    mensAlert(mensText)
                else:
                    st.session_state['containers'] = True
                    st.session_state['allUpLoads'] = uploadedFiles
                    st.session_state['bytesAll'] = fullFiles(uploadedFiles, valKeyPill)
            case _:
                st.session_state['uploaderKey'] += 1            
        scroll_to_element("keyPill")
        st.session_state['keyPill'] = None
        if optSel == 6: 
            st.session_state['disabSlid'] = True
            st.session_state['disabPillTwo'] = True
        else:
            st.session_state['disabSlid'] = False
    elif mode == 1:
        value = st.session_state['numSldImg']
        if value == 0:
            keySel = "keyPill"
        else:
            value -= 1   
        listKeyImgs = list(st.session_state['bytesAll'].keys())
        keyValue = listKeyImgs[value]
        keySel = st.session_state['bytesAll'][keyValue][0][2]
        scroll_to_element(keySel)
        st.session_state['numSldImg'] = 0
    return  

def pillConfigExib(uploadedFiles, optSel):
    if optSel == 0:
        mensInfoAll()
    else:
        mensCreateAll(uploadedFiles)
        
def compareLoads():
    keysData = list(st.session_state['bytesAll'].keys())
    angles = [st.session_state['bytesAll'][key][0][3] for key in keysData]
    return angles

def pillFuncSave(uploadedFiles):
    keysData = list(st.session_state['bytesAll'].keys())
    funcSel = st.session_state['keyPillFive']
    funcNum = optFunc.index(funcSel)
    anglesOne = compareLoads()
    st.session_state['containers'] = True
    st.session_state['allUpLoads'] = uploadedFiles
    st.session_state['bytesAll'] = fullFiles(uploadedFiles, funcSel)
    keysData = list(st.session_state['bytesAll'].keys())
    anglesTwo = compareLoads()
    clickAng = st.session_state['clickAngle']
    if clickAng:
        for a, ang in enumerate(anglesOne): 
            angOne = ang
            angTwo = anglesTwo[a]
            if angOne != angTwo: 
                key = keysData[a]
                st.session_state['bytesAll'][key][0][3] = angOne            
    if funcNum == 0:
        saveMultImgPdf()  
    else:
        saveMultImgDocx()
    st.session_state['keyPillFive'] = None
    st.session_state['disabSlid'] = False
    st.session_state['clickAngle'] = False
       
def main():
    global optFiles, buttSymbs, optPages
    global optAll, optText, optInfo, optCreate 
    global scopeText, optFunc, optAllPapers
    global optAllOrients, optAllMargins
    global optAllAngles, optLabesAll
    optAllOrients = ["paisagem", "retrato"]
    optAllMargins = {'grande': 0.80, 'máxima': 0.75, 'média': 0.85, 'mínima': 0.95, 'pequena': 0.90}
    optAllAngles = [angle for angle in range(-360, 420, 90)]
    optFiles, buttSymbs, optAll, optText, optInfo, optCreate, optPages, scopeText, optFunc, optAllPapers = setVars()
    setPage()
    setSession()
    nCols = 4
    buttKeys = list(buttSymbs.keys())
    nButts = len(buttSymbs)
    buttSufix = [buttSymbs[key][1] for key in buttKeys]
    nameKey = lambda a, b: f'{a}_{str(b+1).zfill(3)}' 
    exts = ['BMP', 'GIF', 'ICO', 'JPG', 'PNG', 'PPM', 'TIF']
    textHead = ':material/app_registration: <span style="color: #73505B;">Mescla de imagens e criação de arquivo pdf ou docx</span>'
    with st.container(border=True, vertical_alignment="top", horizontal_alignment="center"):
        colPills, colConfig = st.columns([15, 5.5])
        colUpload, colPage = st.columns([15, 5.5], vertical_alignment="center")
        uploadedFiles = colUpload.file_uploader(label="Escolha seu arquivo", type=exts, key=f"uploader_{st.session_state['uploaderKey']}", 
                                                accept_multiple_files=True, label_visibility="collapsed", 
                                                on_change=zeraVal, max_upload_size=20*1024)
        nUps = len(uploadedFiles)
        checkDown(uploadedFiles)
        with colConfig:
            colPillInfo, colPillCreate = st.columns(2, vertical_alignment="center", width="stretch", 
                                                    gap="xxsmall") 
            colPillInfo.pills(label="information", options=optInfo, label_visibility="collapsed", 
                              key='keyPillFour', width="stretch", on_change=pillConfigExib, args=(uploadedFiles, 0))
            colPillCreate.pills(label="create", options=optCreate, label_visibility="collapsed", 
                                key='keyPillThree', disabled=st.session_state['disabPillThree'], 
                                width="stretch", on_change=pillConfigExib, args=(uploadedFiles, 1))                     
        colPills.pills(label="options", options=optFiles, label_visibility="collapsed", 
                       on_change=changePill, args=(uploadedFiles, 0), key="keyPill", 
                       disabled=st.session_state['disabPill'], width="stretch") 
        if any([nUps == 0, uploadedFiles is None]):
            nUps = 1
            numStr = 1
        else:
            numStr = f"1 a {nUps}"
        with colPage:
            st.space("xxsmall")
            pillFunc = st.pills(label="funções", options=optFunc, label_visibility="collapsed", on_change=pillFuncSave, 
                                args=(uploadedFiles, ), width="stretch", key='keyPillFive', disabled=st.session_state['disabPillFive'])
            colOne, colTwo = st.columns([10, 3.5], vertical_alignment="center", width="stretch", gap="xxsmall")
            colOne.slider(label="Informe o número da página", min_value=0, max_value=nUps, key="numSldImg",
                          label_visibility="collapsed", width="stretch", on_change=fromToImg, 
                          disabled=st.session_state['disabSlid'])
            colTwo.pills(label="pages", options=optPages, label_visibility="collapsed", on_change=changePill, 
                         args=(None, 1), key='keyPillTwo', disabled=st.session_state['disabPillTwo'], width="stretch")
    if st.session_state['containers']:
        with st.container(border=False, vertical_alignment="top", horizontal_alignment="center"):
            keysData = list(st.session_state['bytesAll'].keys())
            nBts = len(keysData)
            limK = nBts - 1
            for k, key in enumerate(keysData):
                bytesData = st.session_state['bytesAll'][key]
                for bt, btData in enumerate(bytesData):
                    nameFile = btData[0]
                    bytesDataRed = btData[1]
                    keyFile = btData[2]
                    angleFile = btData[3]
                    funcFile = btData[4]
                    with st.container(border=True, vertical_alignment="top", horizontal_alignment="center"):
                        textImg = designateImgs(nameFile, k, nBts, angleFile, funcFile, 0)
                        colBadMark, colMark = st.columns([0.3, 10], vertical_alignment="center", width="stretch", 
                                                          gap="xxsmall")
                        colBadMark.badge(":material/tag:", color="red", width="content")
                        colMark.markdown(textImg, width="stretch")
                        colsButt = st.columns(spec=nButts, width="stretch", gap="small", 
                                              vertical_alignment="top")
                        for b, butt in enumerate(buttKeys):
                            ico, key = buttSymbs[butt]
                            newKeyFile = f'{nameKey(f'{keyFile}_{key}', k)}'
                            buttDisab = False
                            if k == 0 and b in [0, 1]: 
                                buttDisab = True
                            if k == limK and b in [2, 3]:
                                buttDisab = True
                            if b == 0:
                                keyOne = keyFile
                                if k == 0:
                                    st.session_state['keyFirst'] = keyFile
                            else:
                                keyOne = newKeyFile
                            if b in [1, 2]:
                                argsFunc = (newKeyFile, buttSufix, k)
                            else:
                                argsFunc = (newKeyFile, buttSufix, btData, k)
                            colsButt[b].button(label = f'{butt}', key=keyOne, icon=ico, use_container_width=True, 
                                               on_click=commandButt, args=argsFunc, disabled=buttDisab)
                        st.space("xxsmall")
                        st.image(bytesDataRed, caption=nameFile)
                    st.space("xxsmall")
                st.session_state['keyLast'] = newKeyFile 
            if st.session_state['numSlidesAll'] != 0:
                mensSliderAll()
                    
def checkDown(uploadedFiles):
    nUps = len(uploadedFiles)
    if all([nUps > 0, uploadedFiles is not None]):
        st.session_state['disabPill'] = False 
        st.session_state['disabPillThree'] = False
        st.session_state['disabPillFive'] = False
    else:
        st.session_state['containers'] = False
        st.session_state['bytesAll'].clear()
        st.session_state['disabPill'] = True
        st.session_state['disabPillThree'] = True
        st.session_state['disabPillFive'] = True
        st.session_state['allUpLoads'].clear()

def setPage():
    st.set_page_config(
        page_title='Mescla de imagens',
        page_icon=':material/image:',
        layout='wide', 
        initial_sidebar_state=None, 
        menu_items=None)
    #fileCss = r'C:\Users\ACER\Documents\css\configImg_new.css' #(local)
    fileCss = 'configImg_new.css' #(github)
    with open(fileCss) as f:
        css = f.read()
    st.markdown(f'<style>{css}</style>', unsafe_allow_html=True)
 
def setSession():
    keyVals = {'containers': False, 'uploaderKey': 0, 'numSldImg': 0, 
               'bytesAll': {}, 'allUpLoads': [], 'keyLast': '', 'keyFirst': '', 'keyPill': None, 
               'keyPillTwo': None, 'keyPillThree': None, 'angleRotated': optAllAngles[4], 'angleRotatedAll': optAllAngles[4], 
               'disabSlid': True, 'disabPill': True, 'disabPillTwo': True, 
               'disabPillThree': True, 'numSlidesAll': 0.0, 'disabPillFive': True, 
               'keyPillFive': None, 'numResolAll': 200, 'papelSelAll': list(optAllPapers.keys())[6], 
               'orientSelAll': optAllOrients[1], 'marginSelAll': list(optAllMargins.keys())[-2], 'checkYes': False, 
               'checkNo': False, 'clickAngle': False}
    for key, val in keyVals.items():
        if key not in st.session_state:
            st.session_state[key] = val

def setVars():
    optFiles = [":material/order_approve:", ":material/edit_arrow_up:", 
                ":material/edit_arrow_down:", ":material/tile_large:", 
                ":material/tile_small:", ":material/slideshow:", 
                ":material/cleaning_services:"]
    optPages = [":material/find_in_page:"]
    optInfo = [":material/help_clinic:"]
    optCreate = [":material/factory:"]
    optFunc = [":material/save:", ":material/file_save:"]
    optAll = optFiles + optInfo + optCreate + optPages + optFunc
    optText = ["Sequência por ordem normal de upload", 
               "Sequência por ordem alfabética crescente", 
               "Sequência por ordem alfabética decrescente", 
               "Sequência por ordem crescente de tamanho", 
               "Sequência por ordem decrescente de tamanho", 
               "Exibição de fotos como slides", 
               "Limpeza de dados e objetos mostrados na tela", 
               "Informações e detalhamento sobre o app", 
               "Configuração das imagens", 
               "Direcionamento à imagem de número especificado", 
               "Criação de arquivo Pdf", "Criação de arquivo Docx"]
    buttSymbs = {'trás': [':material/chevron_backward:', 'info_info'], 
                 'topo': [':material/first_page:', 'top_top'], 
                 'final': [':material/last_page:', 'end_end'], 
                 'frente': [':material/chevron_forward:', 'forw_forw'],
                 'rotação': [':material/cameraswitch:', 'byte_byte'],
                }
    optCommand = ["Recuo para a imagem anterior", 
                  "Recuo para o topo da tela", "Avanço para o final da tela",
                  "Avanço para a imagem seguinte", "Rotação da imagem"]
    optText += optCommand
    nComm = len(optCommand)
    info = lambda w: buttSymbs[w] 
    optAll += [f"{info(w)[0]} {w}" for w in list(buttSymbs.keys())] 
    nText = len(optText) - nComm
    scopeText = ['Todos os arquivos' for k in range(nText)] 
    scopeText += ["Imagem selecionada" for k in range(nComm)]
    optAllPapers = papersize.SIZES
    return(optFiles, buttSymbs, optAll, optText, optInfo, optCreate, optPages, scopeText, optFunc, optAllPapers)

if __name__ == '__main__':
    try:
        main()
    except Exception as fail:
        mensText = f":blue[**:material/error:**] Houve o seguinte erro {fail}. Contate o administrador."
        mensError(mensText)
    finally:
        st.cache_data.clear()
